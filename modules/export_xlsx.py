import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


_HEADER_FILL  = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT  = Font(color="FFFFFF", bold=True, size=9)
_SUBHEADER_FILL = PatternFill("solid", fgColor="2E75B6")
_SUBHEADER_FONT = Font(color="FFFFFF", bold=True, size=9)
_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)
_SI_FILL  = PatternFill("solid", fgColor="C6EFCE")
_NO_FILL  = PatternFill("solid", fgColor="FFCCCC")


def _style_cell(cell, fill=None, font=None, bold=False, center=False):
    cell.border = _BORDER
    if fill:  cell.fill = fill
    if font:  cell.font = font
    elif bold: cell.font = Font(bold=True, size=9)
    else:      cell.font = Font(size=9)
    if center: cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    else:      cell.alignment = Alignment(vertical="center", wrap_text=True)


def export_gerencia_xlsx(full_df: pd.DataFrame, gerencia: str, periodo: str) -> bytes:
    df = full_df[full_df["gerencia"] == gerencia].copy().reset_index(drop=True)
    si  = (df["subio_ficha"] == "SI").sum()
    no  = len(df) - si
    pct = round(si / len(df) * 100, 1) if len(df) else 0

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = gerencia[:31]

    # Resumen superior
    ws.merge_cells("G1:I1"); ws["G1"] = "CANTIDAD"; ws["G1"].font = Font(bold=True, size=9)
    ws["G2"] = "SI";  ws["H2"] = int(si)
    ws["G3"] = "NO";  ws["H3"] = int(no)
    ws["G4"] = "TOTAL"; ws["H4"] = len(df); ws["I4"] = f"{pct}%"

    # Encabezado
    headers = ["CÓDIGO","APELLIDOS Y NOMBRES","GERENCIA",
               "GERENCIA / SUBGERENCIA / JEFATURA","CARGO",
               "SUBIÓ FICHA\nSI / NO","COMENTARIO",
               "NRO DE OBJETIVOS\nASIGNADOS","FORMATO FIRMADO\nSI / NO"]
    col_widths = [12, 35, 25, 35, 25, 12, 40, 12, 14]

    header_row = 6
    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=header_row, column=col, value=h)
        _style_cell(cell, fill=_HEADER_FILL, font=_HEADER_FONT, center=True)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[header_row].height = 30

    # Datos
    row_map = ["codigo","nombre","gerencia","subgerencia","cargo",
               "subio_ficha","comentario","nro_objetivos","formato_firmado"]
    for r_idx, row in df.iterrows():
        excel_row = header_row + 1 + r_idx
        for c_idx, field in enumerate(row_map, 1):
            val = row.get(field)
            if pd.isna(val): val = ""
            cell = ws.cell(row=excel_row, column=c_idx, value=val)
            _style_cell(cell, center=(c_idx in [1, 6, 8, 9]))
            if field == "subio_ficha":
                cell.fill = _SI_FILL if val == "SI" else _NO_FILL
            if field == "formato_firmado":
                cell.fill = _SI_FILL if val == "SI" else _NO_FILL

    ws.freeze_panes = f"A{header_row + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_all_xlsx(full_df: pd.DataFrame, periodo: str) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    summary = []
    for gerencia in sorted(full_df["gerencia"].dropna().unique()):
        df = full_df[full_df["gerencia"] == gerencia]
        si  = (df["subio_ficha"] == "SI").sum()
        pct = round(si / len(df) * 100, 1) if len(df) else 0
        summary.append({"GERENCIA": gerencia, "REQUERIDAS": len(df),
                        "REGISTRADAS": int(si), "% AVANCE": f"{pct}%"})

        ws = wb.create_sheet(title=gerencia[:31])
        headers = ["CÓDIGO","APELLIDOS Y NOMBRES","GERENCIA",
                   "GERENCIA / SUBGERENCIA / JEFATURA","CARGO",
                   "SUBIÓ FICHA","COMENTARIO","NRO OBJETIVOS","FORMATO FIRMADO"]
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            _style_cell(cell, fill=_HEADER_FILL, font=_HEADER_FONT, center=True)
        row_map = ["codigo","nombre","gerencia","subgerencia","cargo",
                   "subio_ficha","comentario","nro_objetivos","formato_firmado"]
        for r_idx, row in df.reset_index(drop=True).iterrows():
            for c_idx, field in enumerate(row_map, 1):
                val = row.get(field); val = "" if pd.isna(val) else val
                cell = ws.cell(row=r_idx+2, column=c_idx, value=val)
                _style_cell(cell)
                if field in ("subio_ficha","formato_firmado"):
                    cell.fill = _SI_FILL if val == "SI" else _NO_FILL

    # Hoja REPORTE
    ws_r = wb.create_sheet(title="REPORTE", index=0)
    ws_r["A1"] = f"REPORTE FICHAS DE OBJETIVOS INDIVIDUALES — PERIODO {periodo}"
    ws_r["A1"].font = Font(bold=True, size=11)
    for c, h in enumerate(["GERENCIA","REQUERIDAS","REGISTRADAS","% AVANCE"], 1):
        cell = ws_r.cell(row=3, column=c, value=h)
        _style_cell(cell, fill=_HEADER_FILL, font=_HEADER_FONT, center=True)
    for r, row in enumerate(summary, 4):
        for c, val in enumerate(row.values(), 1):
            _style_cell(ws_r.cell(row=r, column=c, value=val))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
