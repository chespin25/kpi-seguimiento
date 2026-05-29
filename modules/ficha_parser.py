import os
import re
import unicodedata
import openpyxl
from config import FICHAS_DIR, FICHA_SHEET, FICHA_COL_N


_CODE_RE  = re.compile(r"^(\d{7})")
_NAME_KW  = re.compile(r"(APELLIDO|NOMBRE|TRABAJADOR|SERVIDOR)", re.IGNORECASE)


def _norm(s: str) -> str:
    return unicodedata.normalize("NFD", s.upper()).encode("ascii", "ignore").decode().strip()


def _extract_code(filename: str) -> str | None:
    m = _CODE_RE.match(os.path.basename(filename))
    return m.group(1) if m else None


def _extract_name_from_xlsx(filepath: str) -> str | None:
    """
    Abre el xlsx y busca el nombre del trabajador en las primeras 30 filas.
    Estrategia: celda con keyword NOMBRE/APELLIDO → devuelve celda adyacente.
    """
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        sheet = wb[FICHA_SHEET] if FICHA_SHEET in wb.sheetnames else wb.active
        rows = list(sheet.iter_rows(max_row=30, values_only=True))
        for row in rows:
            for j, cell in enumerate(row):
                if cell is None:
                    continue
                text = str(cell).strip()
                if _NAME_KW.search(text) and len(text) < 60:
                    # buscar valor en la misma fila (cols siguientes)
                    for k in range(j + 1, min(j + 6, len(row))):
                        val = row[k]
                        if val and str(val).strip() and not _NAME_KW.search(str(val)):
                            candidate = str(val).strip()
                            if len(candidate) > 5 and re.search(r"[A-Za-zÁáÉéÍíÓóÚúÑñ]", candidate):
                                return candidate
        return None
    except Exception:
        return None


_STOP_GER = {"de", "del", "la", "el", "los", "las", "y", "e", "o", "u", "en", "a", "por"}

def _build_name_map(workers_df) -> dict:
    """Retorna {nombre_normalizado: codigo} desde workers_df."""
    name_map = {}
    for _, row in workers_df.iterrows():
        cod = str(row.get("codigo", "") or "").strip()
        nom = str(row.get("nombre", "") or "").strip()
        if cod and nom:
            name_map[_norm(nom)] = cod
    return name_map


def _build_gerentes_map(workers_df) -> dict:
    """
    Retorna {norm_gerencia: codigo_gerente}.
    Infiere la gerencia desde el CARGO del tipo 'GERENTE X':
    'GERENTE LEGAL' → LEGAL, 'GERENTE ADMINISTRACIÓN Y LOGÍSTICA' → ADM...
    Esto evita errores por encargos que mueven al gerente a otra gerencia.
    """
    gerencias = list(workers_df["gerencia"].dropna().unique())
    gerentes: dict = {}

    for _, row in workers_df.iterrows():
        cargo = str(row.get("cargo", "") or "").strip().upper()
        cod   = str(row.get("codigo", "") or "").strip()
        if not cargo.startswith("GERENTE") or not cod:
            continue
        # Extraer la parte tras "GERENTE [DE] "
        after = re.sub(r"^GERENTE\s*(DE\s+)?", "", cargo).strip()
        if not after:
            continue
        # Buscar gerencia con mayor solapamiento de tokens significativos
        after_tokens = {t for t in _norm(after).split()
                        if len(t) > 2 and t.lower() not in _STOP_GER}
        best_ger, best_score = None, 0.0
        for ger in gerencias:
            ger_tokens = {t for t in _norm(ger).split()
                          if len(t) > 2 and t.lower() not in _STOP_GER}
            if not ger_tokens:
                continue
            score = len(ger_tokens & after_tokens) / len(ger_tokens)
            if score >= 0.5 and score > best_score:
                best_score, best_ger = score, ger
        if best_ger:
            key = _norm(best_ger)
            if key not in gerentes:
                gerentes[key] = cod
    return gerentes


def _match_filename_to_gerencia(filename: str, gerencias: list[str]) -> str | None:
    """
    Intenta asociar un filename sin código a una gerencia por similitud de tokens.
    Retorna el nombre de gerencia si hay match suficiente, o None.
    """
    stem = re.sub(r"[_\-]", " ", os.path.splitext(os.path.basename(filename))[0])
    stem_norm = _norm(stem)
    stem_tokens = {t for t in stem_norm.split() if len(t) > 2 and t.lower() not in _STOP_GER}

    best_ger, best_score = None, 0
    for ger in gerencias:
        ger_norm   = _norm(ger)
        ger_tokens = {t for t in ger_norm.split() if len(t) > 2 and t.lower() not in _STOP_GER}
        if not ger_tokens:
            continue
        overlap = len(ger_tokens & stem_tokens)
        score   = overlap / len(ger_tokens)   # fracción de tokens de gerencia presentes
        if score >= 0.5 and overlap >= 1 and score > best_score:
            best_score, best_ger = score, ger
    return best_ger


def _match_extracted_name(extracted: str, name_map: dict) -> str | None:
    """
    Compara el nombre extraído del xlsx contra name_map.
    Acepta: coincidencia exacta normalizada, o todos los tokens del nombre
    de planta presentes en el texto extraído.
    """
    norm_extracted = _norm(extracted)
    # Exacto
    if norm_extracted in name_map:
        return name_map[norm_extracted]
    # Tokens: todos los tokens del nombre de planta deben estar en el extraído
    ext_tokens = set(norm_extracted.split())
    best_cod, best_score = None, 0
    for nom_norm, cod in name_map.items():
        tokens = set(nom_norm.split())
        if len(tokens) < 2:
            continue
        matches = tokens & ext_tokens
        if matches == tokens:          # match completo
            if len(tokens) > best_score:
                best_score, best_cod = len(tokens), cod
    return best_cod


def _detect_signed(filepath: str) -> str:
    """
    Detecta si la ficha tiene formato firmado.
    SI si: hay imagen insertada (firma escaneada) o celda 'FIRM*' con valor adyacente.
    """
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        sheet = wb[FICHA_SHEET] if FICHA_SHEET in wb.sheetnames else wb.active
        # Imagen insertada → firma escaneada
        if getattr(sheet, "_images", None):
            return "SI"
        # Buscar celda con 'FIRM' y valor adyacente no vacío
        for row in sheet.iter_rows(max_row=60, values_only=True):
            for j, cell in enumerate(row):
                if cell is None:
                    continue
                if "FIRM" in str(cell).upper():
                    for k in range(j + 1, min(j + 6, len(row))):
                        v = row[k]
                        if v and str(v).strip() not in ("", "None"):
                            return "SI"
        return "NO"
    except Exception:
        return "NO"


def _count_objectives(filepath: str) -> int | None:
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        if FICHA_SHEET not in wb.sheetnames:
            return None
        ws = wb[FICHA_SHEET]

        col_idx = None
        header_row = None
        for row in ws.iter_rows(max_row=20, values_only=True):
            for i, cell in enumerate(row):
                if str(cell).strip() == FICHA_COL_N:
                    col_idx = i
                    header_row = row
                    break
            if col_idx is not None:
                break

        if col_idx is None:
            return None

        count = 0
        for row in ws.iter_rows(min_row=(ws.min_row or 1), values_only=True):
            if row is header_row:
                continue
            val = row[col_idx] if col_idx < len(row) else None
            if val is not None:
                try:
                    int(float(str(val)))
                    count += 1
                except (ValueError, TypeError):
                    pass
        return count if count > 0 else None
    except Exception:
        return None


def scan_fichas(periodo: str, base_dir: str = None, workers_df=None) -> list[dict]:
    """
    Escanea base_dir (o FICHAS_DIR/<periodo>/) y retorna lista de dicts:
    {"codigo": ..., "subio_ficha": "SI", "nro_objetivos": N o None}
    Para xlsx sin código de 7 dígitos, intenta extraer el nombre desde dentro
    del archivo y matchearlo contra workers_df.
    """
    folder = base_dir or os.path.join(FICHAS_DIR, periodo)
    if not os.path.isdir(folder):
        return []

    name_map     = {}
    gerentes_map = {}
    gerencias_list: list[str] = []
    if workers_df is not None and not workers_df.empty:
        name_map      = _build_name_map(workers_df)
        gerentes_map  = _build_gerentes_map(workers_df)
        gerencias_list = list(workers_df["gerencia"].dropna().unique())

    results = []
    seen = set()
    for fname in os.listdir(folder):
        ext = os.path.splitext(fname)[1].lower()
        if ext not in (".xlsx", ".xls", ".pdf"):
            continue

        fpath  = os.path.join(folder, fname)
        codigo = _extract_code(fname)

        if not codigo and ext in (".xlsx", ".xls"):
            # 1. Intentar extraer nombre desde dentro del xlsx
            if name_map:
                extracted = _extract_name_from_xlsx(fpath)
                if extracted:
                    codigo = _match_extracted_name(extracted, name_map)

            # 2. Si sigue sin código, intentar match por gerencia → asignar al gerente
            if not codigo and gerentes_map and gerencias_list:
                ger = _match_filename_to_gerencia(fname, gerencias_list)
                if ger:
                    codigo = gerentes_map.get(_norm(ger))

        if not codigo or codigo in seen:
            continue
        seen.add(codigo)

        entry = {"codigo": codigo, "subio_ficha": "SI", "nro_objetivos": None, "formato_firmado": "NO"}
        if ext in (".xlsx", ".xls"):
            entry["nro_objetivos"]   = _count_objectives(fpath)
            entry["formato_firmado"] = _detect_signed(fpath)
        results.append(entry)

    return results
